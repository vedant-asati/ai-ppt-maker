"""
pptx_exporter.py — Export presentation as an editable PPTX file.

Strategy (Tier 2 — Visible Overlays):
  Each slide has:
    1. Full-bleed PNG background  — exact visual from the HTML→PNG pipeline
    2. Visible, editable text overlays — layout-aware, theme-colored, proper font sizes

  Layout-specific positioning ensures text boxes sit where they appear visually,
  so the user can: edit a title, fix a bullet, change a number — without
  touching the background image for small tweaks.

  Text colors are derived from the theme's --text CSS variable (already resolved),
  so dark themes get white text, light themes get dark text automatically.

Usage:
    from pptx_exporter import PptxExporter
    exporter = PptxExporter()
    exporter.export(png_files, master_plan, "output.pptx", resolved_vars=vars)
"""

import os
from typing import List, Optional, Dict, Tuple
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from rich import print as rprint


# ─── Slide dimensions (16:9) ─────────────────────────────────────────────────

SLIDE_W = Inches(13.333)   # 1920px at 144dpi
SLIDE_H = Inches(7.5)      # 1080px at 144dpi


# ─── Layout-specific text box positions ───────────────────────────────────────
# Format: (left, top, width, height) in Inches
# Designed around 13.333" × 7.5" slide canvas

LAYOUT_POSITIONS: Dict[str, Dict[str, Tuple]] = {
    'cover': {
        'eyebrow':   (2.0, 2.0,  9.3, 0.5),
        'title':     (1.5, 2.6,  10.3, 1.6),
        'subtitle':  (2.0, 4.3,  9.3, 0.8),
        'content':   (2.0, 5.3,  9.3, 1.8),
    },
    'split': {
        'eyebrow':   (0.8, 0.5,  5.8, 0.45),
        'title':     (0.8, 1.0,  5.8, 1.3),
        'subtitle':  (0.8, 2.4,  5.8, 0.65),
        'content':   (0.8, 3.2,  5.8, 3.8),
    },
    'stat': {
        'eyebrow':   (2.0, 0.5,  9.3, 0.5),
        'title':     (0.8, 1.1,  11.7, 0.9),   # title above the big number
        'value':     (1.0, 2.1,  11.3, 2.0),
        'label':     (2.0, 4.2,  9.3, 0.7),
        'content':   (2.0, 5.1,  9.3, 2.0),
    },
    'quote': {
        'eyebrow':   (1.5, 0.8,  10.3, 0.45),
        'quote':     (1.5, 1.8,  10.3, 3.5),
        'attribution': (2.5, 5.5, 8.3, 0.7),
    },
    'generic': {
        'eyebrow':   (0.8, 0.5,  11.7, 0.45),
        'title':     (0.8, 1.0,  11.7, 1.1),
        'subtitle':  (0.8, 2.2,  11.7, 0.65),
        'content':   (0.8, 3.0,  11.7, 4.0),
    },
    'timeline': {
        'eyebrow':   (0.8, 0.5,  11.7, 0.45),
        'title':     (0.8, 1.0,  11.7, 1.1),
        'subtitle':  (0.8, 2.2,  11.7, 0.65),
        'content':   (0.8, 3.0,  11.7, 4.0),
    },
    'infographic': {
        'eyebrow':   (0.8, 0.5,  11.7, 0.45),
        'title':     (0.8, 1.0,  11.7, 1.1),
        'subtitle':  (0.8, 2.2,  11.7, 0.65),
        'content':   (0.8, 3.0,  11.7, 4.0),
    },
}


# ─── Font size presets ────────────────────────────────────────────────────────

FONT_SIZES = {
    'cover_title':  Pt(36),
    'cover_sub':    Pt(20),
    'eyebrow':      Pt(12),
    'title':        Pt(28),
    'subtitle':     Pt(16),
    'stat_value':   Pt(72),
    'stat_label':   Pt(18),
    'quote_text':   Pt(26),
    'attribution':  Pt(14),
    'bullet':       Pt(16),
    'content':      Pt(15),
}


class PptxExporter:
    """Export rendered slides + metadata to an editable PPTX file."""

    def export(
        self,
        png_files: List[str],
        master_plan,
        output_path: str,
        resolved_vars: Optional[Dict] = None,
    ) -> bool:
        """
        Create a PPTX file with rendered PNGs as backgrounds
        and visible, layout-aware editable text overlays.

        Args:
            png_files:     Ordered list of rendered slide PNG paths.
            master_plan:   MasterPlan with slide metadata.
            output_path:   Path for the output .pptx file.
            resolved_vars: Theme CSS variable dict (for text/accent colors).
        """
        rprint("PptxExporter: Creating editable PPTX...")

        if not png_files:
            rprint("  [Warning] PptxExporter: no PNG files — cannot produce PPTX.")
            return False

        # Derive text and accent colors from theme vars
        text_color, accent_color = self._resolve_colors(resolved_vars or {})

        prs = Presentation()
        prs.slide_width = SLIDE_W
        prs.slide_height = SLIDE_H

        blank_layout = prs.slide_layouts[6]   # Blank — no placeholders

        slides_map = {s.slide_number: s for s in master_plan.slides}

        for i, png_path in enumerate(png_files):
            slide_number = i + 1
            slide_meta = slides_map.get(slide_number)

            if not os.path.exists(png_path):
                rprint(f"  [Warning] Missing PNG for slide {slide_number}: {png_path}")
                continue

            pptx_slide = prs.slides.add_slide(blank_layout)

            # ── Full-bleed background PNG ───────────────────────────────────
            pptx_slide.shapes.add_picture(
                png_path, Emu(0), Emu(0), SLIDE_W, SLIDE_H,
            )

            # ── Speaker notes ───────────────────────────────────────────────
            if slide_meta and slide_meta.speaker_notes:
                notes = pptx_slide.notes_slide
                notes.notes_text_frame.text = slide_meta.speaker_notes

            # ── Visible, layout-aware text overlays ─────────────────────────
            if slide_meta:
                self._add_overlays(pptx_slide, slide_meta, text_color, accent_color)

        # ── Save ────────────────────────────────────────────────────────────
        out_dir = os.path.dirname(output_path)
        if out_dir:
            os.makedirs(out_dir, exist_ok=True)

        try:
            prs.save(output_path)
            rprint(f"PptxExporter: Saved to {output_path}")
            return True
        except PermissionError:
            alt = output_path.replace(".pptx", "_v2.pptx")
            rprint(f"  [Warning] {output_path} locked — saving as {alt}")
            prs.save(alt)
            return True
        except Exception as e:
            rprint(f"  [Error] PPTX export failed: {e}")
            return False

    # ── Color resolution ──────────────────────────────────────────────────────

    def _resolve_colors(
        self, resolved_vars: Dict
    ) -> Tuple[RGBColor, RGBColor]:
        """
        Derive text + accent colors from theme CSS vars.

        --text  → body text color (white for dark themes, dark for light themes)
        --primary → accent color (used for stat values, eyebrow labels)
        Falls back to white text / gold accent if vars are missing.
        """
        def hex_to_rgb(hex_val: str) -> RGBColor:
            h = hex_val.lstrip('#')
            if len(h) != 6:
                raise ValueError(f"Bad hex: {hex_val}")
            return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))

        try:
            text_color = hex_to_rgb(resolved_vars.get('--text', '#FFFFFF'))
        except (ValueError, TypeError):
            text_color = RGBColor(255, 255, 255)

        try:
            accent_color = hex_to_rgb(resolved_vars.get('--primary', '#F5C518'))
        except (ValueError, TypeError):
            accent_color = RGBColor(245, 197, 24)

        return text_color, accent_color

    # ── Layout-aware overlays ─────────────────────────────────────────────────

    def _add_overlays(self, pptx_slide, slide_meta, text_color: RGBColor, accent_color: RGBColor):
        """Add visible, positioned text overlays based on layout type."""
        layout = slide_meta.layout_type or 'generic'
        positions = LAYOUT_POSITIONS.get(layout, LAYOUT_POSITIONS['generic'])

        is_centered = layout in ('cover', 'stat')
        title_align = PP_ALIGN.CENTER if is_centered else PP_ALIGN.LEFT

        # ── Eyebrow ──────────────────────────────────────────────────────
        if slide_meta.eyebrow and 'eyebrow' in positions:
            self._textbox(pptx_slide, slide_meta.eyebrow.upper(),
                          positions['eyebrow'], FONT_SIZES['eyebrow'],
                          accent_color, bold=True, alignment=title_align)

        # ── Stat layout ──────────────────────────────────────────────────
        if layout == 'stat':
            self._add_stat_overlays(pptx_slide, slide_meta, positions,
                                    text_color, accent_color)

        # ── Quote layout ─────────────────────────────────────────────────
        elif layout == 'quote':
            self._add_quote_overlays(pptx_slide, slide_meta, positions, text_color)

        # ── Cover / Split / Generic / Timeline / Infographic ─────────────
        else:
            # Title
            if slide_meta.title and 'title' in positions:
                title_size = (FONT_SIZES['cover_title'] if layout == 'cover'
                              else FONT_SIZES['title'])
                self._textbox(pptx_slide, slide_meta.title,
                              positions['title'], title_size,
                              text_color, bold=True, alignment=title_align)

            # Subtitle
            if slide_meta.subtitle and 'subtitle' in positions:
                sub_size = (FONT_SIZES['cover_sub'] if layout == 'cover'
                            else FONT_SIZES['subtitle'])
                self._textbox(pptx_slide, slide_meta.subtitle,
                              positions['subtitle'], sub_size,
                              text_color, italic=True, alignment=title_align)

            # Content blocks
            if 'content' in positions:
                content_lines = self._extract_content(slide_meta)
                if content_lines:
                    self._textbox(pptx_slide, '\n'.join(content_lines),
                                  positions['content'], FONT_SIZES['content'],
                                  text_color, alignment=PP_ALIGN.LEFT,
                                  wrap=True)

    def _add_stat_overlays(self, pptx_slide, slide_meta, positions,
                           text_color: RGBColor, accent_color: RGBColor):
        """Special handling for stat layout — big number + label."""
        # Title above the stat
        if slide_meta.title and 'title' in positions:
            self._textbox(pptx_slide, slide_meta.title,
                          positions['title'], FONT_SIZES['title'],
                          text_color, bold=True, alignment=PP_ALIGN.CENTER)

        # Extract stat block
        stat_block = None
        if slide_meta.content_blocks:
            for block in slide_meta.content_blocks:
                if block.block_type == 'stat':
                    stat_block = block
                    break

        if stat_block and 'value' in positions:
            self._textbox(pptx_slide, stat_block.value or '',
                          positions['value'], FONT_SIZES['stat_value'],
                          accent_color, bold=True, alignment=PP_ALIGN.CENTER)
            if stat_block.label and 'label' in positions:
                self._textbox(pptx_slide, stat_block.label,
                              positions['label'], FONT_SIZES['stat_label'],
                              text_color, alignment=PP_ALIGN.CENTER)
        elif 'content' in positions:
            # Fallback: no stat block found — render as generic content
            content_lines = self._extract_content(slide_meta)
            if content_lines:
                self._textbox(pptx_slide, '\n'.join(content_lines),
                              positions['content'], FONT_SIZES['content'],
                              text_color, wrap=True)

    def _add_quote_overlays(self, pptx_slide, slide_meta, positions, text_color: RGBColor):
        """Special handling for quote layout — centered italic quote + attribution."""
        quote_block = None
        if slide_meta.content_blocks:
            for block in slide_meta.content_blocks:
                if block.block_type == 'pullquote':
                    quote_block = block
                    break

        if quote_block and 'quote' in positions:
            quote_text = f'"{quote_block.quote}"' if quote_block.quote else ''
            self._textbox(pptx_slide, quote_text,
                          positions['quote'], FONT_SIZES['quote_text'],
                          text_color, italic=True,
                          alignment=PP_ALIGN.CENTER, wrap=True)
            if quote_block.attribution and 'attribution' in positions:
                self._textbox(pptx_slide, f"— {quote_block.attribution}",
                              positions['attribution'], FONT_SIZES['attribution'],
                              text_color, italic=True, alignment=PP_ALIGN.CENTER)
        elif slide_meta.title and 'quote' in positions:
            # Fallback: use title as quote text
            self._textbox(pptx_slide, f'"{slide_meta.title}"',
                          positions['quote'], FONT_SIZES['quote_text'],
                          text_color, italic=True, alignment=PP_ALIGN.CENTER, wrap=True)

    # ── Content extraction ────────────────────────────────────────────────────

    def _extract_content(self, slide_meta) -> List[str]:
        """Flatten content blocks into readable lines for a text overlay."""
        lines = []

        for block in (slide_meta.content_blocks or []):
            bt = block.block_type

            if bt == 'bullets' and block.items:
                for item in block.items:
                    lines.append(f"• {item}")

            elif bt == 'stat':
                lines.append(f"{block.value}  —  {block.label}")

            elif bt == 'pullquote' and block.quote:
                lines.append(f'"{block.quote}"')
                if block.attribution:
                    lines.append(f"— {block.attribution}")

            elif bt == 'steps' and block.steps:
                labels = block.step_labels or []
                for j, step in enumerate(block.steps):
                    label = labels[j] if j < len(labels) else f"Step {j+1}"
                    lines.append(f"{label}: {step}")

            elif bt == 'keyvalue' and block.pairs:
                for pair in block.pairs:
                    k = pair.get('key', '')
                    v = pair.get('value', '')
                    lines.append(f"{k}: {v}")

            elif bt == 'comparison':
                l_label = block.left_label or 'Option A'
                r_label = block.right_label or 'Option B'
                l_items = ', '.join(block.left_items or [])
                r_items = ', '.join(block.right_items or [])
                lines.append(f"{l_label}: {l_items}")
                lines.append(f"{r_label}: {r_items}")

            elif bt == 'callout' and block.text:
                lines.append(f"💡 {block.text}")

        # Legacy fallback
        if not lines and slide_meta.content_bullets:
            lines = [f"• {b}" for b in slide_meta.content_bullets]

        return lines

    # ── Text box primitive ────────────────────────────────────────────────────

    def _textbox(
        self,
        slide,
        text: str,
        position: Tuple[float, float, float, float],
        font_size,
        color: RGBColor,
        bold: bool = False,
        italic: bool = False,
        alignment: PP_ALIGN = PP_ALIGN.LEFT,
        wrap: bool = True,
    ):
        """
        Add a single positioned, formatted text box to a slide.

        Text is always visible (proper color + size).
        Font: Calibri — clean, universally available in PowerPoint.
        """
        if not text or not text.strip():
            return

        left, top, width, height = [Inches(v) for v in position]

        txBox = slide.shapes.add_textbox(left, top, width, height)
        tf = txBox.text_frame
        tf.word_wrap = wrap
        tf.auto_size = None   # Don't auto-resize — respect our position bounds

        # First paragraph
        p = tf.paragraphs[0]
        p.alignment = alignment

        # Handle multi-line text (newlines → separate paragraphs)
        lines = text.split('\n')
        for i, line in enumerate(lines):
            if i == 0:
                para = p
            else:
                para = tf.add_paragraph()
                para.alignment = alignment

            run = para.add_run()
            run.text = line
            run.font.name = 'Calibri'
            run.font.size = font_size
            run.font.bold = bold
            run.font.italic = italic
            run.font.color.rgb = color
